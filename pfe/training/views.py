from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated, BasePermission, AllowAny
from rest_framework.response import Response
from .models import *
from .serializers import *
from rest_framework import status
from users.models import *
from django.utils import timezone
from datetime import timedelta
from openpyxl import Workbook
from django.http import HttpResponse
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from users.views import *
from django.db.models import Count

@api_view(['POST'])
@permission_classes([AllowAny])#IsEmployeur
def post_training_needs(request):
    serializer= TrainingNeedSerializer(data=request.data, partial=True)
    if serializer.is_valid():
        form = TrainingForm.objects.filter(manager=request.user).order_by('-created_at').first()
        if form.status=='SENT':
            return Response({"message": "can't add more needs, the form is already submited!"},status=400)        
        serializer.save(created_by=request.user,form=form)
        return Response(serializer.data,status=status.HTTP_201_CREATED)
    return Response(serializer.errors,status=400)
    
@api_view(['GET','DELETE','PATCH'])
@permission_classes([AllowAny])#IsEmployeur
def update_training_need(request,pk):
    try:
        need=TrainingNeed.objects.get(pk=pk)
        form=need.form
        if form.status=='SENT':
            return Response({"message": "form already submited"},status=400)
    except TrainingNeed.DoesNotExist:
       return Response({"meesage": "training need not found"}, status=400)
    if request.method=="PATCH":
        serializer= TrainingNeedSerializer(need, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data,status=201)
        return Response(serializer.errors,status=400)
    elif request.method=="DELETE":
        need.delete()
        return Response({"message": "training need deleted"})
    elif request.method=="GET":
        serializer=TrainingNeedSerializer(need)
        return Response(serializer.data,status=200)
        
@api_view(['GET'])
@permission_classes([AllowAny])#IsAuthenticated
def get_trainings(request):
    trainings=Training.objects.all()
    serializer=TrainingSerializer(trainings,many=True)
    return Response(serializer.data,status=200)

# @api_view(['GET','POST','DELETE'])
# def training_response(request,pk):
#     try:
#         training=Training.objects.get(pk=pk)
#     except Training.DoesNotExist:
#         return Response({"message": "training does not exist!"},status=400)
#     if request.method=='GET':
#         serializer=TrainingSerializer(training)
#         return Response(serializer.data,status=200)
#     elif request.method=='PATCH':
#         serializer=TrainingSerializer(training,data=request.data,partial=True)
#         if serializer.is_valid():
#             serializer.save()
#             return Response(serializer.data,status=200)
#         return Response(serializer.errors,status=400)
#     elif request.type=='DELETE':
#         training.delete()
#         return Response({"message": "training deleted!"},status=200)
    
@api_view(['POST'])
@permission_classes([AllowAny])#IsDDRH
def make_decesion(request,pk):
    try:
        need=TrainingNeed.objects.get(pk=pk)
        form=need.form
    except TrainingNeed.DoesNotExist:
        return Response({"message": "training need does not exist!"},status=400)
    serializer=DecisionSerializers(data=request.data,partial=True)
    if serializer.is_valid():
        decesion=serializer.save(TrainingNeed=need, user=request.user)
        need.status=decesion.result
        need.save()
    
        submited_needs=TrainingNeed.objects.filter(form=form)
        
        if not    submited_needs.filter(status='WAITING').exists():
            form.status='HANDLED'
            form.save()
        if decesion.result=='APPROVED':
            training=Training.objects.create(
                title=need.title,
                TrainingNeed=need,
                status='notassigned',
                type=request.data.get('type')
                )
            training.save()
        return Response({"message:" "decesion add status changed"},status=200)
        
    return Response(serializer.errors,status=400)

@api_view(['POST','GET'])
@permission_classes([AllowAny])#IsDDRH
def get_post_forms(request):
    if request.method=='GET':
        users=User.objects.filter(role='MANAGER')
        serializer=UserSerializer(users,many=True)
        return Response(serializer.data)    
    # 9olo l rayan ydir confirmation!!
    elif request.method=='POST':
        print(request.data)
        manager_ids = request.data.get('managers')
        if  not manager_ids:
            return Response({"message": "managers are required"}, status=400)
        created_forms = []
        for manager_id in manager_ids:
            try:
                manager = User.objects.get(id=manager_id, role='MANAGER')
            except User.DoesNotExist:
                continue
            form = TrainingForm.objects.create(
                manager=manager
            )
            created_forms.append(form.id)
            notification=Notification.objects.create(
                user=manager,
                title='Training form',
                message='make sure to fill the training form on the link down below!'
            )
        return Response({
            "message": "forms sent",
            "forms": created_forms
        })
    #token: eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJ0b2tlbl90eXBlIjoiYWNjZXNzIiwiZXhwIjoxNzc3OTI1MDg3LCJpYXQiOjE3Nzc5MjQ3ODcsImp0aSI6IjdmOTc3MGI4ZjJhYjRhZDBhNDAyM2YzZmYxZTE4MzgyIiwidXNlcl9pZCI6IjExIn0.iFP3Co4upsbruPBNLCjSQLvuDTzDGFmzV1HxYgXj5oU
@api_view(['GET','POST'])
@permission_classes([AllowAny])#IsEmployeur
def manager_form(request):
    user=request.user
    if request.method=='GET':
        form=TrainingForm.objects.latest("created_at")
        if timezone.now().date() > form.finalDate:
            return Response({"message": "form is closed!"},status=401)
        forms=TrainingForm.objects.filter(created_at=form.created_at)
        try:
            currentForm=forms.get(manager=user)
        except TrainingForm.DoesNotExist:
            return Response({"message": "you dont have access to the form"},status=401)
        
        needs=TrainingNeed.objects.filter(created_by=user,form=currentForm)
        serializer=TrainingNeedSerializer(needs,many=True)
        return Response({"needs":serializer.data})
    elif request.method=='POST':
        form=TrainingForm.objects.filter(manager=user).order_by('-created_at').first()
        form.status='SENT'
        form.save()
        return Response({"message": "form submited"})        
        
        
        
        
        
        
@api_view(['GET'])
@permission_classes([AllowAny])#IsDDRH
def Generate_report(request):
    
    year = request.GET.get('year')

    if not year:
        return Response({"error": "year is required"}, status=400)
    decisions = Decision.objects.filter(date__year=year)\
        .select_related('TrainingNeed', 'TrainingNeed__created_by', 'user')
    wb = Workbook()
    ws = wb.active
    ws.title = f"Decisions {year}"
    headers = [
        "Training Need",
        "Manager",
        "Result",
        "Decision By",
        "Comment",
        "Date"
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(size=16, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    for d in decisions:
        ws.append([
            d.TrainingNeed.title,
            d.TrainingNeed.created_by.username,  
            d.result,
            d.user.username,
            d.comment,
            d.date.strftime("%Y-%m-%d")
        ])

    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)

        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

            cell.alignment = Alignment(wrap_text=True)

        ws.column_dimensions[col_letter].width = min(max_length + 2, 40)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=decision_report_{year}.xlsx'

    wb.save(response)
    return response

@api_view(['GET'])
@permission_classes([AllowAny])#IsDDRH
def access_submited_forms(request):
    forms=TrainingForm.objects.filter(status='SENT')
    serializer=FormSerializer(forms,many=True)
    return Response({"forms": serializer.data })
@api_view(['GET'])
@permission_classes([AllowAny])
def manage_submited_forms(request,pk):
    try:
        form=TrainingForm.objects.get(pk=pk)
    except TrainingForm.DoesNotExist:
        return Response({"message": "form does not exist"},status=400)
    needs=TrainingNeed.objects.filter(form=form)
    serializer=TrainingNeedSerializer(needs,many=True)
    return Response({"needs": serializer.data})

@api_view(['GET'])
@permission_classes([AllowAny])
def ddrh_dashboard(request):
    fiches_envoyees = TrainingForm.objects.count()
    formations_demandees = TrainingNeed.objects.count()
    formations_validees = TrainingNeed.objects.filter(status='APPROVED').count()
    formations_refusees = TrainingNeed.objects.filter(status='DENIED').count()

    # Fiches à valider = besoins pas encore décidés
    fiches_a_valider = TrainingNeed.objects.exclude(
        status__in=['APPROVED', 'DENIED']
    ).count()

    # Utilisateurs managers / employeurs
    utilisateurs = User.objects.filter(role='MANAGER').count()

    # État des besoins / fiches pour donut chart
    etat_fiches = (
        TrainingNeed.objects
        .values('status')
        .annotate(total=Count('id'))
        .order_by('status')
    )

    # Formations les plus demandées
    formations_plus_demandees = (
        TrainingNeed.objects
        .values('title')
        .annotate(total=Count('id'))
        .order_by('-total')[:5]
    )

    data = {
        "stats": {
            "fiches_envoyees": fiches_envoyees,
            "fiches_a_valider": fiches_a_valider,
            "formations_demandees": formations_demandees,
            "formations_validees": formations_validees,
            "formations_refusees": formations_refusees,
            "utilisateurs": utilisateurs,
        },
        "etat_fiches": list(etat_fiches),
        "formations_plus_demandees": list(formations_plus_demandees),
    }

    return Response(data, status=200)
    